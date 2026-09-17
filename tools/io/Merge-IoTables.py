# -*- coding: utf-8 -*-
"""Parse matrix IO Excel + existing TIA Hardware.xml, emit merge plan and updated XML."""
from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
LEGACY = Path(r"C:\Users\Hank\Documents\Codex\2026-08-14\new-chat-3\outputs\TiaOpennessCheck")
ROOT = LEGACY  # Practice/HmiExport 仍在舊包；下一波再遷
OUT = ROOT / "Practice" / "io-merge"
C_XLSX = Path(r"C:\Users\Hank\Desktop\25017\複製25017ST_大亞_630mm-63B框絞機產線_新.xlsx")
B_XLSX = Path(r"C:\Users\Hank\Desktop\24137ST_拉梅莎_630mm-37B 框絞機生產線.xlsx")
HW_ROOT = ROOT / "HmiExport" / "Templates" / "26037_Cusor"

# C sheet -> TIA PLC software name (post-rename 25017_*)
SHEET_TO_PLC = {
    "LINE PLC": ["25017_Main_PLC"],
    "TF-260 PLC": ["25017_TF_PLC"],
    "Rotor inside PLC": [
        "25017_6B_Inside_PLC",
        "25017_12B_Inside_PLC",
        "25017_18B_Inside_PLC",
    ],
}

EXPORT_FOLDER = {
    "25017_Main_PLC": "PLC_24137_Main_PLC",
    "25017_TF_PLC": "PLC_24137_TF_PLC",
    "25017_6B_Inside_PLC": "PLC_24137_6B_Inside_PLC",
    "25017_12B_Inside_PLC": "PLC_24137_12B_Inside_PLC",
    "25017_18B_Inside_PLC": "PLC_24137_18B_Inside_PLC",
}

ADDR_CELL = re.compile(
    r"^(?:%?\s*)(I|Q|IW|QW|ID|QD|AIW|AQW)(\d+(?:\.\d+)?)$",
    re.I,
)
NOISE = {
    "hub", "x1", "p1", "p2", "no.", "no", "spare", "di x 14 (dc)", "di x 16 (dc)",
    "dq x 10 (dc)", "dq x 16 (dc)", "sm 1238 energy meter",
}


def cell_str(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        return ""
    s = str(v).replace("\n", " ").strip()
    return s


def norm_addr(raw: str) -> str | None:
    m = ADDR_CELL.match(raw.replace(" ", ""))
    if not m:
        return None
    kind, num = m.group(1).upper(), m.group(2)
    if kind == "AIW":
        kind = "IW"
    if kind == "AQW":
        kind = "QW"
    return f"%{kind}{num}"


def datatype_of(addr: str) -> str:
    body = addr[1:]
    if body.startswith("IW") or body.startswith("QW"):
        return "Int"
    if body.startswith("ID") or body.startswith("QD"):
        return "DInt"
    return "Bool"


def looks_desc(s: str) -> bool:
    if not s or len(s) < 2:
        return False
    low = s.lower()
    if low in NOISE:
        return False
    if ADDR_CELL.match(s.replace(" ", "")):
        return False
    if re.fullmatch(r"\d+(\.\d+)?", s):
        return False
    return True


def parse_matrix_sheet(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    points = []
    seen = set()
    for ri, row in enumerate(rows, 1):
        if not row:
            continue
        for ci, v in enumerate(row):
            addr = norm_addr(cell_str(v))
            if not addr:
                continue
            desc = ""
            if ci + 1 < len(row):
                nxt = cell_str(row[ci + 1])
                if looks_desc(nxt):
                    desc = nxt
            key = (addr, desc)
            if key in seen:
                continue
            seen.add(key)
            # skip duplicate address with empty desc if we already have one with desc
            points.append({"row": ri, "col": ci + 1, "address": addr, "desc": desc})
    # collapse same address: keep longest desc
    by_addr: dict[str, dict] = {}
    for p in points:
        prev = by_addr.get(p["address"])
        if prev is None or len(p["desc"]) > len(prev["desc"]):
            by_addr[p["address"]] = p
    return [by_addr[k] for k in sorted(by_addr, key=addr_sort_key)]


def addr_sort_key(a: str):
    m = re.match(r"%([A-Z]+)(\d+)(?:\.(\d+))?", a)
    if not m:
        return (9, a)
    kind = m.group(1)
    order = {"I": 0, "Q": 1, "IW": 2, "QW": 3, "ID": 4, "QD": 5}.get(kind, 8)
    bit = int(m.group(3) or 0)
    return (order, int(m.group(2)), bit)


def parse_xlsx(path: Path, sheets: list[str] | None = None) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    want = sheets or list(wb.sheetnames)
    out = {}
    for sn in want:
        if sn not in wb.sheetnames:
            continue
        out[sn] = parse_matrix_sheet(wb[sn])
    wb.close()
    return out


def parse_hardware_xml(path: Path) -> list[dict]:
    tree = ET.parse(path)
    root = tree.getroot()
    tags = []
    for tag in root.iter():
        if not tag.tag.endswith("PlcTag"):
            continue
        attrs = {}
        comment = ""
        for child in list(tag):
            local = child.tag.split("}")[-1]
            if local == "AttributeList":
                for a in list(child):
                    attrs[a.tag.split("}")[-1]] = (a.text or "").strip()
            elif local == "ObjectList":
                for text_el in child.iter():
                    if text_el.tag.split("}")[-1] == "Text" and (text_el.text or "").strip():
                        comment = text_el.text.strip()
                        break
        addr = attrs.get("LogicalAddress", "")
        name = attrs.get("Name", "")
        if not addr or not name:
            continue
        tags.append({
            "name": name,
            "address": addr,
            "type": attrs.get("DataTypeName", "Bool"),
            "comment": comment,
        })
    return tags


def slug_tag(addr: str, desc: str) -> str:
    prefix = "I" if addr.startswith("%I") or addr.startswith("%IW") or addr.startswith("%ID") else "Q"
    body = addr[1:].replace(".", "_")
    spare = bool(re.search(r"spare|備用|预留|預留", desc, re.I))
    if spare or not desc:
        return f"{prefix}_{body}"
    # keep address-based unique names for new points
    return f"{prefix}_{body}"


def hex_id(n: int) -> str:
    return format(n, "X")


def make_tag_xml(tag_id: int, name: str, addr: str, dtype: str, comment: str) -> ET.Element:
    tag = ET.Element("SW.Tags.PlcTag", ID=hex_id(tag_id), CompositionName="Tags")
    al = ET.SubElement(tag, "AttributeList")
    ET.SubElement(al, "DataTypeName").text = dtype
    ET.SubElement(al, "ExternalAccessible").text = "true"
    ET.SubElement(al, "ExternalVisible").text = "true"
    ET.SubElement(al, "ExternalWritable").text = "true"
    ET.SubElement(al, "LogicalAddress").text = addr
    ET.SubElement(al, "Name").text = name
    ol = ET.SubElement(tag, "ObjectList")
    mt = ET.SubElement(ol, "MultilingualText", ID=hex_id(tag_id + 1), CompositionName="Comment")
    items = ET.SubElement(mt, "ObjectList")
    cultures = [("en-US", comment), ("zh-CN", comment), ("zh-TW", comment)]
    for i, (cult, text) in enumerate(cultures):
        item = ET.SubElement(items, "MultilingualTextItem", ID=hex_id(tag_id + 2 + i), CompositionName="Items")
        ial = ET.SubElement(item, "AttributeList")
        ET.SubElement(ial, "Culture").text = cult
        ET.SubElement(ial, "Text").text = text
    return tag


def rebuild_hardware_xml(src: Path, merged: list[dict], dest: Path) -> None:
    tree = ET.parse(src)
    root = tree.getroot()
    table = None
    for el in root.iter():
        if el.tag.endswith("PlcTagTable"):
            table = el
            break
    if table is None:
        raise RuntimeError("no PlcTagTable in " + str(src))
    obj = None
    for child in list(table):
        if child.tag.endswith("ObjectList"):
            obj = child
            break
    if obj is None:
        obj = ET.SubElement(table, "ObjectList")
    for child in list(obj):
        obj.remove(child)
    tid = 1
    for t in merged:
        node = make_tag_xml(tid, t["name"], t["address"], t["type"], t.get("comment") or t.get("desc") or "")
        obj.append(node)
        tid += 6
    dest.parent.mkdir(parents=True, exist_ok=True)
    tree.write(dest, encoding="utf-8", xml_declaration=True)


def merge_one(plc: str, excel_pts: list[dict], hw: list[dict]) -> dict:
    hw_by_addr = {t["address"]: t for t in hw}
    used_names = {t["name"] for t in hw}
    excel_by_addr = {p["address"]: p for p in excel_pts}

    updated, created, unchanged = [], [], []
    merged = []

    for t in hw:
        xp = excel_by_addr.get(t["address"])
        item = dict(t)
        if xp and xp["desc"] and xp["desc"] != t["comment"]:
            item["comment"] = xp["desc"]
            item["action"] = "update-comment"
            item["excel_desc"] = xp["desc"]
            updated.append(item)
        else:
            item["action"] = "keep"
            unchanged.append(item)
        merged.append(item)

    for p in excel_pts:
        if p["address"] in hw_by_addr:
            continue
        name = slug_tag(p["address"], p["desc"])
        base = name
        n = 2
        while name in used_names:
            name = f"{base}_{n}"
            n += 1
        used_names.add(name)
        item = {
            "name": name,
            "address": p["address"],
            "type": datatype_of(p["address"]),
            "comment": p["desc"],
            "action": "create",
        }
        created.append(item)
        merged.append(item)

    only_hw = [t for t in hw if t["address"] not in excel_by_addr]
    return {
        "plc": plc,
        "excel": len(excel_pts),
        "hardware": len(hw),
        "updated": updated,
        "created": created,
        "unchanged": unchanged,
        "only_in_hardware": only_hw,
        "merged": merged,
    }


def write_csv(path: Path, rows: list[dict], cols: list[str]) -> None:
    import csv
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    c_sheets = parse_xlsx(C_XLSX, list(SHEET_TO_PLC) + ["PF-260 PLC", "Rotor outside PLC"])
    b_sheets = {}
    if B_XLSX.exists():
        b_sheets = parse_xlsx(B_XLSX, ["LINE PLC", "Rotor inside PLC", "TF PLC"])

    summary = []
    all_created = []
    all_updated = []

    for sheet, plcs in SHEET_TO_PLC.items():
        pts = c_sheets.get(sheet, [])
        for plc in plcs:
            folder = EXPORT_FOLDER[plc]
            hw_path = HW_ROOT / folder / "TagTables" / "Hardware.xml"
            hw = parse_hardware_xml(hw_path) if hw_path.exists() else []
            result = merge_one(plc, pts, hw)
            summary.append({
                "sheet": sheet,
                "plc": plc,
                "excel_pts": result["excel"],
                "hardware_pts": result["hardware"],
                "update_comment": len(result["updated"]),
                "create": len(result["created"]),
                "keep": len(result["unchanged"]),
                "only_hw": len(result["only_in_hardware"]),
            })
            for x in result["created"]:
                all_created.append({"plc": plc, "sheet": sheet, **x})
            for x in result["updated"]:
                all_updated.append({"plc": plc, "sheet": sheet, "name": x["name"],
                                    "address": x["address"], "old": x.get("comment"),
                                    "new": x.get("excel_desc")})

            dest = OUT / f"PLC_{plc}" / "TagTables" / "Hardware.xml"
            rebuild_hardware_xml(hw_path, result["merged"], dest)

            write_csv(OUT / f"diff_{plc}.csv",
                      result["merged"] + [
                          {**t, "action": "only-hardware"} for t in result["only_in_hardware"]
                          if t["address"] not in {m["address"] for m in result["merged"]}
                      ],
                      ["action", "name", "address", "type", "comment"])

    write_csv(OUT / "summary.csv", summary,
              ["sheet", "plc", "excel_pts", "hardware_pts", "update_comment", "create", "keep", "only_hw"])
    write_csv(OUT / "created.csv", all_created,
              ["plc", "sheet", "name", "address", "type", "comment"])
    write_csv(OUT / "updated.csv", all_updated,
              ["plc", "sheet", "name", "address", "old", "new"])

    extra = {}
    for sn in ("PF-260 PLC", "Rotor outside PLC"):
        extra[sn] = [{"address": p["address"], "desc": p["desc"]} for p in c_sheets.get(sn, [])]
    (OUT / "unmapped-sheets.json").write_text(
        json.dumps({k: {"count": len(v), "sample": v[:8]} for k, v in extra.items()},
                   ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (OUT / "excel-counts.json").write_text(
        json.dumps({
            "C": {k: len(v) for k, v in c_sheets.items()},
            "B": {k: len(v) for k, v in b_sheets.items()},
        }, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print("WROTE", OUT)
    for row in summary:
        print("{plc}: excel={excel_pts} hw={hardware_pts} upd={update_comment} new={create} keep={keep} onlyHw={only_hw}".format(**row))


if __name__ == "__main__":
    main()
