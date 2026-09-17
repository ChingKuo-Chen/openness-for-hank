# -*- coding: utf-8 -*-
"""Build full Hardware.xml from 25017 IO table + existing tags; write problem list."""
from __future__ import annotations

import csv
import re
import xml.etree.ElementTree as ET
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
LEGACY = Path(r"C:\Users\Hank\Documents\Codex\2026-08-14\new-chat-3\outputs\TiaOpennessCheck")
ROOT = LEGACY  # Practice/HmiExport 仍在舊包；下一波再遷
OUT = ROOT / "Practice" / "io-merge"
IMPORT = OUT / "import"
C_XLSX = next(Path(r"C:\Users\Hank\Desktop\25017").glob("*25017ST*63B*.xlsx"))
HW_ROOT = ROOT / "HmiExport" / "Templates" / "26037_Cusor"
THRESHOLD = 0.80

PLC_MAP = {
    "LINE PLC": [
        "25017_Main_PLC",
        "PLC_24137_Main_PLC",
    ],
    "TF-260 PLC": [
        "25017_TF_PLC",
        "PLC_24137_TF_PLC",
    ],
    "Rotor inside PLC": [
        "25017_6B_Inside_PLC",
        "PLC_24137_6B_Inside_PLC",
    ],
}
# extra inside PLCs share the same C sheet
INSIDE_EXTRA = [
    ("25017_12B_Inside_PLC", "PLC_24137_12B_Inside_PLC"),
    ("25017_18B_Inside_PLC", "PLC_24137_18B_Inside_PLC"),
]

ADDR_CELL = re.compile(r"^(?:%?\s*)(I|Q|IW|QW|ID|QD|AIW|AQW)(\d+(?:\.\d+)?)$", re.I)
STRIP = re.compile(
    r"LINE\s*|全線\s*|PBL|PB\s*|OP\d(?:/\d)*|JOYSTICK|\(NO\)|\(NC\)|\(RLY\)|\(MC\)|"
    r"\s+|　|[()（）,_、\-–]",
    re.I,
)
OP_RE = re.compile(r"OP\s*([12])", re.I)

NAME_HINTS = [
    (r"計尺.*CH-A|譯碼器.*CH-A", "I_Enc_Length_A"),
    (r"計尺.*CH-B|譯碼器.*CH-B", "I_Enc_Length_B"),
    (r"給線架伺服.*CH-A", "I_Enc_1B_A"),
    (r"給線架伺服.*CH-B", "I_Enc_1B_B"),
    (r"6B.*歸零", "I_Snr_6B_Home"),
    (r"12B.*歸零", "I_Snr_12B_Home"),
    (r"20B.*歸零", "I_Snr_20B_Home"),
    (r"24B.*歸零", "I_Snr_24B_Home"),
    (r"引取馬達超溫\s*-B", "I_F_Cap_Mtr_B_OvrHt2"),
    (r"引取馬達過載\s*-B", "I_F_Cap_Mtr_B_Ovrld2"),
    (r"捲取驅動器三相電源", "I_F_Takeup_Power_Mon"),
    (r"緊急停止.*OP1", "I_PB_EStop_OP1"),
    (r"電鈴按鈕.*OP2", "I_PB_Bell_OP2"),
    (r"捲取啟動.*OP2", "I_PB_Takeup_Start_OP2"),
    (r"捲取停止.*OP2", "I_PB_Takeup_Stop_OP2"),
    (r"緊急停止.*OP2", "I_PB_EStop_OP2"),
    (r"正寸動", "I_PB_Cap_Jog_Fwd"),
    (r"逆寸動", "I_PB_Cap_Jog_Rev"),
    (r"壓線按鈕", "I_PB_Cap_Clamp"),
    (r"引取.*緊急停止", "I_PB_EStop_Cap"),
    (r"捲取_運轉$", "I_S_Takeup_Run"),
    (r"UPS 狀態.*1", "I_S_UPS_1"),
    (r"UPS 狀態.*2", "I_S_UPS_2"),
    (r"UPS 狀態.*3", "I_S_UPS_3"),
    (r"20B轉體_DC BUS", "Q_20B_DC_Bus_On"),
    (r"24B轉體_馬達冷卻風扇", "Q_24B_Mtr_Fan"),
    (r"24B轉體_DC BUS", "Q_24B_DC_Bus_On"),
    (r"故障燈 PBL \(OP1\)", "Q_Lamp_Alarm_OP1"),
    (r"捲取運轉燈 PBL \(OP1\)", "Q_Lamp_Takeup_Run_OP1"),
    (r"捲取停止燈 PBL \(OP1\)", "Q_Lamp_Takeup_Stop_OP1"),
    (r"捲取停止燈 PBL \(OP2\)", "Q_Lamp_Takeup_Stop_OP2"),
    (r"給線伺服驅動器啟動", "Q_1B_Drv_On"),
    (r"線繃緊|TIGHT", "Q_1B_Tension_Tighten"),
    (r"給線伺服驅動器反向", "Q_1B_Tension_Inverse"),
    (r"給線伺服驅動器緊急停止", "Q_1B_EStop"),
    (r"給線伺服驅動器故障復歸", "Q_1B_Reset"),
    (r"給線_運轉", "Q_Payoff_Run"),
    (r"給線_$", "Q_Payoff_Aux"),
    (r"捲取_運轉 \(RLY\)", "Q_Takeup_Run"),
    (r"捲取_DC BUS", "Q_Takeup_DC_Bus_On"),
    (r"UPS 接收外部", "Q_UPS_Ext_Ctrl"),
    (r"UPS Remote", "Q_UPS_Remote"),
    (r"欠相", "I_F_Phase_Loss"),
    (r"排線伺服驅動器脈波.*-A", "I_Enc_Trv_A"),
    (r"排線伺服驅動器脈波.*-B", "I_Enc_Trv_B"),
    (r"主馬達超溫", "I_F_Mtr_OvrHt"),
    (r"升/降/頂/退馬達過載", "I_F_Lift_Lock_OvrLd"),
    (r"E-STOP 繼電器緊急停止", "I_F_EStop_Relay"),
    (r"緊急停止訊號.*OP3", "I_PB_EStop_OP3"),
    (r"A側機器倾斜", "I_Snr_Tilt_A"),
    (r"B側機器倾斜", "I_Snr_Tilt_B"),
    (r"快速排線往A", "I_PB_Trv_Jog_To_A"),
    (r"快速排線往B", "I_PB_Trv_Jog_To_B"),
    (r"排線換向", "I_PB_Trv_Dir_Change"),
    (r"B側軸下降", "Q_Lift_Down_B"),
    (r"排線往A側方向指示燈", "Q_Lamp_Trv_To_A"),
    (r"排線往B側方向指示燈", "Q_Lamp_Trv_To_B"),
    (r"排線A側極限設定燈", "Q_Lamp_Trv_Limit_A"),
    (r"排線B側極限設定燈", "Q_Lamp_Trv_Limit_B"),
    (r"E-STOP 繼電器復歸", "Q_EStop_Loop_Reset"),
    (r"張力調整", "IW_Tension_VR"),
    (r"鐵軸確認.*#1|#1 段面 鐵軸", "I_Snr_Bobbin_S1"),
    (r"#2 段面 鐵軸", "I_Snr_Bobbin_S2"),
    (r"#3 段面 鐵軸", "I_Snr_Bobbin_S3"),
    (r"#4 段面 鐵軸", "I_Snr_Bobbin_S4"),
    (r"#4 段面安全板打開檢出", "I_Snr_Safe_Pin_Unlocked_S4"),
    (r"#4 段面安全板關閉檢出", "I_Snr_Safe_Pin_locked_S4"),
    (r"#4 段面安全板打開電磁閥", "Q_Safe_Pin_Unlock_S4"),
    (r"#1 段面頂退軸", "Q_Bobbin_Lock_S1"),
    (r"#2 段面頂退軸", "Q_Bobbin_Lock_S2"),
    (r"#3 段面頂退軸", "Q_Bobbin_Lock_S3"),
    (r"#4 段面頂退軸", "Q_Bobbin_Lock_S4"),
    (r"馬達編碼器脈衝", "I_Enc_Rotor"),
    (r"Line speed To DSP", "Q_LineSpeed_To_DSP"),
    (r"Line run To SERVO", "Q_LineRun_To_Servo"),
]


def cell_str(v) -> str:
    if v is None or hasattr(v, "isoformat"):
        return ""
    return str(v).replace("\n", " ").strip()


def norm_addr(raw: str):
    m = ADDR_CELL.match(raw.replace(" ", ""))
    if not m:
        return None
    kind, num = m.group(1).upper(), m.group(2)
    if kind == "AIW":
        kind = "IW"
    if kind == "AQW":
        kind = "QW"
    return f"%{kind}{num}"


def looks_desc(s: str) -> bool:
    if not s or len(s) < 2:
        return False
    if ADDR_CELL.match(s.replace(" ", "")):
        return False
    return s.lower() not in {"hub", "x1", "p1", "p2", "no.", "no"}


def parse_sheet(ws):
    by = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        for ci, v in enumerate(row):
            addr = norm_addr(cell_str(v))
            if not addr:
                continue
            desc = ""
            if ci + 1 < len(row) and looks_desc(cell_str(row[ci + 1])):
                desc = cell_str(row[ci + 1])
            prev = by.get(addr)
            if prev is None or len(desc) > len(prev):
                by[addr] = desc
    return by


def parse_hw(path: Path):
    tree = ET.parse(path)
    tags = []
    for tag in tree.getroot().iter():
        if not tag.tag.endswith("PlcTag"):
            continue
        attrs, comment = {}, ""
        for child in list(tag):
            local = child.tag.split("}")[-1]
            if local == "AttributeList":
                for a in list(child):
                    attrs[a.tag.split("}")[-1]] = (a.text or "").strip()
            elif local == "ObjectList":
                for text_el in child.iter():
                    if text_el.tag.split("}")[-1] == "Text" and (text_el.text or "").strip():
                        comment = text_el.text.strip()
                        break
        if attrs.get("LogicalAddress") and attrs.get("Name"):
            tags.append({
                "name": attrs["Name"],
                "address": attrs["LogicalAddress"],
                "type": attrs.get("DataTypeName", "Bool"),
                "comment": comment,
            })
    return tags, tree


def fold(s: str) -> str:
    return STRIP.sub("", s or "").casefold()


def ratio(a: str, b: str) -> float:
    fa, fb = fold(a), fold(b)
    if not fa and not fb:
        return 1.0
    if not fa or not fb:
        return 0.0
    return SequenceMatcher(None, fa, fb).ratio()


def is_spare(desc: str) -> bool:
    if not (desc or "").strip():
        return True
    return bool(re.search(r"spare|備用|预留|預留", desc, re.I))


def datatype(addr: str) -> str:
    b = addr[1:]
    if b.startswith("IW") or b.startswith("QW"):
        return "Int"
    if b.startswith("ID") or b.startswith("QD"):
        return "DInt"
    return "Bool"


def addr_key(a: str):
    m = re.match(r"%([A-Z]+)(\d+)(?:\.(\d+))?", a)
    kind = {"I": 0, "Q": 1, "IW": 2, "QW": 3, "ID": 4, "QD": 5}.get(m.group(1) if m else "", 8)
    return (kind, int(m.group(2) if m else 0), int(m.group(3) or 0))


def hint_name(desc: str, addr: str) -> str:
    for pat, name in NAME_HINTS:
        if re.search(pat, desc or "", re.I):
            return name
    prefix = "Q" if addr.startswith("%Q") else "I"
    if addr.startswith("%IW") or addr.startswith("%QW"):
        prefix = addr[1:3]
    body = addr[1:].replace(".", "_")
    if is_spare(desc):
        return f"{prefix}_Spare_{body}"
    op = OP_RE.search(desc or "")
    suffix = f"_OP{op.group(1)}" if op else ""
    return f"{prefix}_{body}{suffix}"


def unique(name: str, used: set) -> str:
    if name not in used:
        return name
    n = 2
    while f"{name}_{n}" in used:
        n += 1
    return f"{name}_{n}"


def hex_id(n: int) -> str:
    return format(n, "X")


def make_tag(tid: int, name: str, addr: str, dtype: str, comment: str) -> ET.Element:
    tag = ET.Element("SW.Tags.PlcTag", ID=hex_id(tid), CompositionName="Tags")
    al = ET.SubElement(tag, "AttributeList")
    ET.SubElement(al, "DataTypeName").text = dtype
    ET.SubElement(al, "ExternalAccessible").text = "true"
    ET.SubElement(al, "ExternalVisible").text = "true"
    ET.SubElement(al, "ExternalWritable").text = "true"
    ET.SubElement(al, "LogicalAddress").text = addr
    ET.SubElement(al, "Name").text = name
    ol = ET.SubElement(tag, "ObjectList")
    mt = ET.SubElement(ol, "MultilingualText", ID=hex_id(tid + 1), CompositionName="Comment")
    items = ET.SubElement(mt, "ObjectList")
    for i, cult in enumerate(("en-US", "zh-CN", "zh-TW")):
        item = ET.SubElement(items, "MultilingualTextItem", ID=hex_id(tid + 2 + i), CompositionName="Items")
        ial = ET.SubElement(item, "AttributeList")
        ET.SubElement(ial, "Culture").text = cult
        ET.SubElement(ial, "Text").text = comment or ""
    return tag


def write_xml(src: Path, tags: list[dict], dest: Path):
    tree = ET.parse(src)
    root = tree.getroot()
    table = next(el for el in root.iter() if el.tag.endswith("PlcTagTable"))
    obj = next(ch for ch in list(table) if ch.tag.endswith("ObjectList"))
    for ch in list(obj):
        obj.remove(ch)
    tid = 1
    for t in tags:
        obj.append(make_tag(tid, t["name"], t["address"], t["type"], t["comment"]))
        tid += 6
    dest.parent.mkdir(parents=True, exist_ok=True)
    tree.write(dest, encoding="utf-8", xml_declaration=True)


def op_core(desc: str) -> str:
    s = re.sub(r"OP\s*[12](?:/OP\s*[12])?", "", desc or "", flags=re.I)
    return fold(s)


def find_op_pairs(tags: list[dict]) -> list[tuple]:
    op1, op2 = {}, {}
    for t in tags:
        m = OP_RE.search(t["comment"] or "")
        if not m:
            continue
        core = op_core(t["comment"])
        if not core:
            continue
        if m.group(1) == "1":
            op1[core] = t
        else:
            op2[core] = t
    pairs = []
    for core, a in op1.items():
        b = op2.get(core)
        if b:
            pairs.append((a, b, core))
    return pairs


def merge_plc(plc: str, export_folder: str, excel: dict) -> dict:
    src = HW_ROOT / export_folder / "TagTables" / "Hardware.xml"
    hw, _ = parse_hw(src)
    used = {t["name"] for t in hw}
    hw_by = {t["address"]: t for t in hw}
    merged = []
    conflicts, same, created, kept_only = [], [], [], []

    for t in hw:
        desc = excel.get(t["address"])
        item = dict(t)
        if desc is None:
            item["comment"] = t["comment"]
            item["kind"] = "tia-only"
            kept_only.append(item)
        else:
            r = ratio(t["comment"], desc)
            item["comment"] = desc
            item["old_comment"] = t["comment"]
            item["ratio"] = round(r, 2)
            if r >= THRESHOLD or (not t["comment"] and not is_spare(desc)):
                item["kind"] = "same"
                same.append(item)
            elif r >= 0.55:
                item["kind"] = "fuzzy"
                same.append(item)
            else:
                item["kind"] = "conflict"
                conflicts.append(item)
        merged.append(item)

    for addr, desc in excel.items():
        if addr in hw_by:
            continue
        name = unique(hint_name(desc, addr), used)
        used.add(name)
        item = {
            "name": name,
            "address": addr,
            "type": datatype(addr),
            "comment": desc,
            "kind": "new",
        }
        created.append(item)
        merged.append(item)

    merged.sort(key=lambda t: addr_key(t["address"]))
    dest = IMPORT / f"PLC_{plc}" / "TagTables" / "Hardware.xml"
    write_xml(src, merged, dest)
    return {
        "plc": plc,
        "merged": merged,
        "conflicts": conflicts,
        "same": same,
        "created": created,
        "kept_only": kept_only,
        "op_pairs": find_op_pairs(merged),
        "src": src,
        "dest": dest,
    }


def main():
    wb = openpyxl.load_workbook(C_XLSX, read_only=True, data_only=True)
    excel_sheets = {sn: parse_sheet(wb[sn]) for sn in wb.sheetnames}
    wb.close()

    results = []
    line = excel_sheets.get("LINE PLC", {})
    tf = excel_sheets.get("TF-260 PLC", {})
    rotor = excel_sheets.get("Rotor inside PLC", {})

    results.append(merge_plc("25017_Main_PLC", "PLC_24137_Main_PLC", line))
    results.append(merge_plc("25017_TF_PLC", "PLC_24137_TF_PLC", tf))
    results.append(merge_plc("25017_6B_Inside_PLC", "PLC_24137_6B_Inside_PLC", rotor))
    results.append(merge_plc("25017_12B_Inside_PLC", "PLC_24137_12B_Inside_PLC", rotor))
    results.append(merge_plc("25017_18B_Inside_PLC", "PLC_24137_18B_Inside_PLC", rotor))

    pf = excel_sheets.get("PF-260 PLC", {})
    pf_export = "PLC_25017_PF_PLC"
    if (HW_ROOT / pf_export / "TagTables" / "Hardware.xml").exists():
        results.append(merge_plc("25017_PF_PLC", pf_export, pf))
        pf_note = None
    else:
        pf_note = f"- PF-260 PLC：{len(pf)} 點，專案還沒這台，沒建。"

    lines = []
    lines.append("# 25017 IO 已全部建進 Hardware — 待程式討論清單")
    lines.append("")
    lines.append("規則：OP1 與 OP2 同一功能，**程式要 OR 併 IO**。Hardware 兩個位址都建。")
    lines.append("同址 Tag 名沿用 24137；註解改成 C 表。衝突不改程式，列在下面慢慢補。")
    lines.append("")

    for r in results:
        lines.append(f"## {r['plc']}")
        lines.append(
            f"- 寫入 {len(r['merged'])} 點（原有沿用 {len(r['same'])+len(r['conflicts'])+len(r['kept_only'])}，新增 {len(r['created'])}）"
        )
        if r["op_pairs"]:
            lines.append("- **OP1/OP2 程式要併：**")
            for a, b, core in r["op_pairs"]:
                lines.append(
                    f"  - `{a['name']}` {a['address']}  OR  `{b['name']}` {b['address']}  （{a['comment'][:40]} / {b['comment'][:40]}）"
                )
        if r["conflicts"]:
            lines.append("- **同址功能對不上（程式仍用舊名，配線已是 C 表）：**")
            for x in r["conflicts"]:
                lines.append(
                    f"  - {x['address']} `{x['name']}`  舊:{x.get('old_comment','')[:32]}  → C:{x['comment'][:32]}  ({x.get('ratio',0):.0%})"
                )
        lines.append("")

    lines.append("## 未對到 PLC")
    if pf_note:
        lines.append(pf_note)
    lines.append("- Rotor outside PLC：不是矩陣 IO 表，沒建。")

    report = OUT / "PROBLEMS.md"
    report.write_text("\n".join(lines), encoding="utf-8")

    with (OUT / "problems.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["plc", "kind", "address", "name", "old", "excel", "ratio"])
        w.writeheader()
        for r in results:
            for a, b, core in r["op_pairs"]:
                w.writerow({"plc": r["plc"], "kind": "op1-op2", "address": f"{a['address']}|{b['address']}",
                            "name": f"{a['name']}|{b['name']}", "old": a["comment"], "excel": b["comment"], "ratio": ""})
            for x in r["conflicts"]:
                w.writerow({"plc": r["plc"], "kind": "conflict", "address": x["address"], "name": x["name"],
                            "old": x.get("old_comment", ""), "excel": x["comment"], "ratio": x.get("ratio", "")})
            for x in r["created"]:
                w.writerow({"plc": r["plc"], "kind": "new", "address": x["address"], "name": x["name"],
                            "old": "", "excel": x["comment"], "ratio": ""})

    print("import", IMPORT)
    print("problems", report)
    for r in results:
        print(r["plc"], "tags", len(r["merged"]), "new", len(r["created"]),
              "conflict", len(r["conflicts"]), "op-pairs", len(r["op_pairs"]))


if __name__ == "__main__":
    main()
