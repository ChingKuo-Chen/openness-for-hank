# -*- coding: utf-8 -*-
"""Scan Siemens-style IO Excel workbooks and extract point rows."""
import json
import re
import sys
from pathlib import Path

import openpyxl

FILES = {
    "A-22021": Path(r"c:\Users\Hank\Desktop\22021ST_大亞_630mm-61B框絞機生產線.xlsx"),
    "B-24137": Path(r"c:\Users\Hank\Desktop\24137ST_拉梅莎_630mm-37B 框絞機生產線.xlsx"),
    "C-25017": Path(r"c:\Users\Hank\Desktop\25017\複製25017ST_大亞_630mm-63B框絞機產線_新.xlsx"),
}

# PLC IO sheets (skip comm/power/motor name sheets)
SKIP_PATTERNS = re.compile(
    r"通信|通訊|DC BUS|POWER|DELTA|電氣|Electrical|Board|NAME|電機|CT |CT_|線徑|Motor|motor",
    re.I,
)

ADDR_RE = re.compile(r"%[IQMTC][\w\.]+", re.I)
TAG_RE = re.compile(r"^[A-Z][A-Za-z0-9_]{2,}$")

HEADER_HINTS = (
    "位址", "地址", "Address", "PLC位址", "I/O", "IO", "符號", "Tag", "名稱", "Name",
    "說明", "備註", "Comment", "Description", "NO.", "編號", "端子", "Terminal",
    "PLC TAG", "Symbol",
)


def cell_str(v):
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        return v.isoformat(sep=" ", timespec="seconds")
    return str(v).strip()


def row_text(row):
    return " | ".join(cell_str(c) for c in row if cell_str(c))


def score_header(row):
    s = 0
    for c in row:
        t = cell_str(c)
        if not t:
            continue
        for h in HEADER_HINTS:
            if h.lower() in t.lower():
                s += 1
    return s


def find_header_row(rows, limit=80):
    best = (0, -1)
    for i, row in enumerate(rows[:limit]):
        sc = score_header(row)
        if sc > best[0]:
            best = (sc, i)
    return best[1] if best[0] >= 2 else -1


def detect_cols(header):
    cols = {}
    for i, c in enumerate(header):
        t = cell_str(c).lower()
        if not t:
            continue
        if any(x in t for x in ("位址", "地址", "address", "plc位址", "io address")):
            cols.setdefault("address", i)
        elif any(x in t for x in ("符號", "tag", "symbol", "plc tag", "plc_tag")):
            cols.setdefault("tag", i)
        elif any(x in t for x in ("名稱", "name")) and "tag" not in cols:
            cols.setdefault("name", i)
        elif any(x in t for x in ("說明", "備註", "comment", "description", "功能")):
            cols.setdefault("desc", i)
        elif t in ("no.", "no", "編號", "#"):
            cols.setdefault("no", i)
        elif any(x in t for x in ("端子", "terminal", "接線")):
            cols.setdefault("terminal", i)
    return cols


def extract_points(sheet_name, rows):
    hdr_idx = find_header_row(rows)
    points = []
    if hdr_idx < 0:
        # fallback: rows containing %I/%Q
        for ri, row in enumerate(rows):
            txt = row_text(row)
            if ADDR_RE.search(txt):
                points.append({
                    "row": ri + 1,
                    "raw": txt[:300],
                    "address": ADDR_RE.search(txt).group(0),
                })
        return {"mode": "addr_scan", "header_row": None, "cols": {}, "points": points[:500]}

    header = rows[hdr_idx]
    cols = detect_cols(header)
    for ri in range(hdr_idx + 1, len(rows)):
        row = rows[ri]
        if not any(cell_str(c) for c in row):
            continue
        txt = row_text(row)
        addr = None
        if "address" in cols and cols["address"] < len(row):
            addr = cell_str(row[cols["address"]])
        if not addr or not ADDR_RE.search(addr):
            m = ADDR_RE.search(txt)
            addr = m.group(0) if m else ""

        tag = ""
        for key in ("tag", "name"):
            if key in cols and cols[key] < len(row):
                tag = cell_str(row[cols[key]])
                if tag:
                    break

        desc = cell_str(row[cols["desc"]]) if "desc" in cols and cols["desc"] < len(row) else ""

        if not addr and not tag:
            continue
        if addr and not ADDR_RE.search(addr) and not tag:
            continue

        points.append({
            "row": ri + 1,
            "no": cell_str(row[cols["no"]]) if "no" in cols and cols["no"] < len(row) else "",
            "address": addr,
            "tag": tag,
            "desc": desc[:120],
        })
    return {"mode": "header", "header_row": hdr_idx + 1, "cols": cols, "points": points}


def analyze_file(label, path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {"label": label, "path": str(path), "sheets": {}}
    for sn in wb.sheetnames:
        if SKIP_PATTERNS.search(sn):
            continue
        ws = wb[sn]
        rows = [r for r in ws.iter_rows(values_only=True)]
        info = extract_points(sn, rows)
        if info["points"]:
            out["sheets"][sn] = {
                "mode": info["mode"],
                "header_row": info["header_row"],
                "cols": info["cols"],
                "count": len(info["points"]),
                "sample": info["points"][:5],
            }
    wb.close()
    return out


def main():
    report = {}
    all_addrs = {}
    for label, path in FILES.items():
        if not path.exists():
            report[label] = {"error": "file not found"}
            continue
        data = analyze_file(label, path)
        report[label] = data
        for sn, sh in data.get("sheets", {}).items():
            for p in sh.get("sample", []):
                a = p.get("address", "")
                if a:
                    all_addrs.setdefault(a, []).append(f"{label}/{sn}")

    out_path = Path(__file__).resolve().parents[2] / "Practice" / "io-table-scan.json"
    out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    # summary stdout (ascii-safe counts)
    print("WROTE", out_path)
    for label, data in report.items():
        if "error" in data:
            print(label, data["error"])
            continue
        total = sum(s["count"] for s in data["sheets"].values())
        print(f"{label}: {len(data['sheets'])} plc sheets, {total} io rows")
        for sn, sh in sorted(data["sheets"].items(), key=lambda x: -x[1]["count"])[:8]:
            print(f"  {sn}: {sh['count']} ({sh['mode']}) cols={sh['cols']}")


if __name__ == "__main__":
    main()
