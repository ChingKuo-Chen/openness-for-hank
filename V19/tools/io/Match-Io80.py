# -*- coding: utf-8 -*-
"""Match C-table descriptions to existing Hardware comments at ~80% similarity."""
from __future__ import annotations

import csv
import re
import xml.etree.ElementTree as ET
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
LEGACY = Path(r"C:\Users\Hank\Documents\Codex\2026-08-14\new-chat-3\outputs\TiaOpennessCheck")
ROOT = LEGACY  # Practice/HmiExport 仍在舊包；下一波再遷
OUT = ROOT / "Practice" / "io-merge"
C_XLSX = Path(r"C:\Users\Hank\Desktop\25017") / "複製25017ST_大亞_630mm-63B框絞機產線_新.xlsx"
HW_ROOT = ROOT / "HmiExport" / "Templates" / "26037_Cusor"
THRESHOLD = 0.80

SHEET_TO_PLC = {
    "LINE PLC": ["25017_Main_PLC"],
    "TF-260 PLC": ["25017_TF_PLC"],
    "Rotor inside PLC": ["25017_6B_Inside_PLC"],
}

EXPORT_FOLDER = {
    "25017_Main_PLC": "PLC_24137_Main_PLC",
    "25017_TF_PLC": "PLC_24137_TF_PLC",
    "25017_6B_Inside_PLC": "PLC_24137_6B_Inside_PLC",
}

ADDR_CELL = re.compile(r"^(?:%?\s*)(I|Q|IW|QW|ID|QD|AIW|AQW)(\d+(?:\.\d+)?)$", re.I)
NOISE = {"hub", "x1", "p1", "p2", "no.", "no"}

# strip filler so LINE/全線, PBL, OP1 don't kill the ratio
STRIP = re.compile(
    r"LINE\s*|全線\s*|PBL|PB\s*|OP\d(?:/\d)*|JOYSTICK|\(NO\)|\(NC\)|\(RLY\)|\(MC\)|"
    r"備用|Spare|預留|预留|\s+|　|[()（）,_、\-–]",
    re.I,
)


def cell_str(v) -> str:
    if v is None or hasattr(v, "isoformat"):
        return ""
    return str(v).replace("\n", " ").strip()


def norm_addr(raw: str):
    m = ADDR_CELL.match(raw.replace(" ", ""))
    if not m:
        return None
    kind, num = m.group(1).upper(), m.group(2)
    if kind == "AIW":
        kind = "IW"
    if kind == "AQW":
        kind = "QW"
    return f"%{kind}{num}"


def looks_desc(s: str) -> bool:
    if not s or len(s) < 2 or s.lower() in NOISE:
        return False
    if ADDR_CELL.match(s.replace(" ", "")):
        return False
    return not re.fullmatch(r"\d+(\.\d+)?", s)


def parse_sheet(ws):
    by = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        for ci, v in enumerate(row):
            addr = norm_addr(cell_str(v))
            if not addr:
                continue
            desc = ""
            if ci + 1 < len(row) and looks_desc(cell_str(row[ci + 1])):
                desc = cell_str(row[ci + 1])
            prev = by.get(addr)
            if prev is None or len(desc) > len(prev):
                by[addr] = desc
    return by


def parse_hw(path: Path):
    tree = ET.parse(path)
    tags = []
    for tag in tree.getroot().iter():
        if not tag.tag.endswith("PlcTag"):
            continue
        attrs, comment = {}, ""
        for child in list(tag):
            local = child.tag.split("}")[-1]
            if local == "AttributeList":
                for a in list(child):
                    attrs[a.tag.split("}")[-1]] = (a.text or "").strip()
            elif local == "ObjectList":
                for text_el in child.iter():
                    if text_el.tag.split("}")[-1] == "Text" and (text_el.text or "").strip():
                        comment = text_el.text.strip()
                        break
        if attrs.get("LogicalAddress") and attrs.get("Name"):
            tags.append({
                "name": attrs["Name"],
                "address": attrs["LogicalAddress"],
                "type": attrs.get("DataTypeName", "Bool"),
                "comment": comment,
            })
    return tags


def fold(s: str) -> str:
    s = STRIP.sub("", s or "")
    return s.casefold()


def ratio(a: str, b: str) -> float:
    fa, fb = fold(a), fold(b)
    if not fa and not fb:
        return 1.0
    if not fa or not fb:
        return 0.0
    return SequenceMatcher(None, fa, fb).ratio()


def is_spare(desc: str) -> bool:
    if not (desc or "").strip():
        return True
    return bool(re.search(r"spare|備用|预留|預留", desc, re.I))


def main():
    wb = openpyxl.load_workbook(C_XLSX, read_only=True, data_only=True)
    rows_out = []
    listing = []
    for sheet, plcs in SHEET_TO_PLC.items():
        excel = parse_sheet(wb[sheet]) if sheet in wb.sheetnames else {}
        plc = plcs[0]
        hw = parse_hw(HW_ROOT / EXPORT_FOLDER[plc] / "TagTables" / "Hardware.xml")
        hw_by = {t["address"]: t for t in hw}

        same, fuzzy, conflict, created, only_hw = [], [], [], [], []
        for t in hw:
            desc = excel.get(t["address"])
            if desc is None:
                only_hw.append(t)
                continue
            r = ratio(t["comment"], desc)
            item = {**t, "excel": desc, "ratio": round(r, 2)}
            if r >= THRESHOLD:
                same.append(item)
            elif r >= 0.55:
                fuzzy.append(item)
            else:
                conflict.append(item)

        for addr, desc in excel.items():
            if addr not in hw_by:
                created.append({"address": addr, "excel": desc, "spare": is_spare(desc)})

        listing.append(f"==== {plc} / {sheet}")
        listing.append(
            f"  同址且說明≥80%：{len(same)}   55–80%近似：{len(fuzzy)}   "
            f"同址但差很多：{len(conflict)}   新增：{len(created)}   僅TIA有：{len(only_hw)}"
        )
        listing.append("-- 視為同一點（≥80%）--")
        for x in same:
            listing.append(f"  {x['address']:8} {x['name']:32} {x['ratio']:.0%}  {x['comment'][:40]}")
        listing.append("-- 近似（55–80%，也當同一點）--")
        for x in fuzzy:
            listing.append(
                f"  {x['address']:8} {x['name']:32} {x['ratio']:.0%}"
                f"  舊:{x['comment'][:28]}  C:{x['excel'][:28]}"
            )
        listing.append("-- 同址差很多（才要討論）--")
        for x in conflict:
            listing.append(
                f"  {x['address']:8} {x['name']:32} {x['ratio']:.0%}"
                f"  舊:{x['comment'][:28]}  C:{x['excel'][:28]}"
            )
        listing.append("-- 新增（C有、Hardware無）有用 --")
        for x in created:
            if not x["spare"]:
                listing.append(f"  {x['address']:8} {x['excel'][:70]}")
        listing.append("-- 新增備用/空白 --")
        for x in created:
            if x["spare"]:
                listing.append(f"  {x['address']:8} {x['excel'][:40]}")
        listing.append("")

        def dump(kind, items):
            for x in items:
                rows_out.append({
                    "plc": plc, "kind": kind, "address": x.get("address"),
                    "name": x.get("name", ""), "ratio": x.get("ratio", ""),
                    "old_comment": x.get("comment", ""), "excel": x.get("excel", x.get("comment", "")),
                })

        dump("same80", same)
        dump("fuzzy55", fuzzy)
        dump("conflict", conflict)
        dump("new", created)
        dump("tia-only", only_hw)

    wb.close()
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "match80-listing.txt").write_text("\n".join(listing), encoding="utf-8")
    with (OUT / "match80.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["plc", "kind", "address", "name", "ratio", "old_comment", "excel"])
        w.writeheader()
        w.writerows(rows_out)
    print("\n".join(listing))


if __name__ == "__main__":
    main()
