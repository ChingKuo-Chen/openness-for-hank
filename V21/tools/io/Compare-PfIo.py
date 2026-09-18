# -*- coding: utf-8 -*-
"""Compare 22021 給線 PLC Excel vs 25017 PF-260 sheet."""
import json
import os
import re
from pathlib import Path
import openpyxl

ADDR = re.compile(r"^(?:%?\s*)(I|Q|IW|QW|ID|QD)(\d+(?:\.\d+)?)$", re.I)

def cell_str(v):
    if v is None or hasattr(v, "isoformat"):
        return ""
    return str(v).replace("\n", " ").strip()

def norm(raw):
    m = ADDR.match(raw.replace(" ", ""))
    if not m:
        return None
    return "%" + m.group(1).upper() + m.group(2)

def parse_matrix(ws):
    by = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        for ci, v in enumerate(row):
            addr = norm(cell_str(v))
            if not addr:
                continue
            desc = ""
            if ci + 1 < len(row):
                nxt = cell_str(row[ci + 1])
                if nxt and not ADDR.match(nxt.replace(" ", "")) and nxt.lower() not in {"hub", "no.", "no"}:
                    desc = nxt
            if addr not in by or len(desc) > len(by[addr]):
                by[addr] = desc
    return by

desktop = Path(r"C:\Users\Hank\Desktop")
a = Path(os.environ["A_XLSX"]) if os.environ.get("A_XLSX") else next(desktop.glob("22021ST*.xlsx"))
c_dir = desktop / "25017"
c = Path(os.environ["C_XLSX"]) if os.environ.get("C_XLSX") else next(
    p for p in c_dir.glob("*.xlsx") if "63B" in p.name and p.name.endswith("新.xlsx")
)
print("A", a)
print("C", c)
print("A exists", a.exists(), "C exists", c.exists())

wb_a = openpyxl.load_workbook(a, read_only=True, data_only=True)
print("A sheets", [s for s in wb_a.sheetnames])
sheet_a = None
for n in wb_a.sheetnames:
    if "給線" in n or "PF" in n.upper() or "Pay" in n:
        sheet_a = n
        break
print("A pf sheet", sheet_a)
pts_a = parse_matrix(wb_a[sheet_a]) if sheet_a else {}
wb_a.close()

wb_c = openpyxl.load_workbook(c, read_only=True, data_only=True)
print("C sheets", wb_c.sheetnames)
pts_c = parse_matrix(wb_c["PF-260 PLC"]) if "PF-260 PLC" in wb_c.sheetnames else {}
wb_c.close()

out = Path(r"C:\Users\Hank\Documents\Codex\2026-08-14\new-chat-3\outputs\TiaOpennessCheck\Practice\io-merge")
out.mkdir(parents=True, exist_ok=True)
rows = []
addrs = sorted(set(pts_a) | set(pts_c))
same = diff = only_a = only_c = 0
lines = ["# 22021 給線 vs 25017 PF-260 IO", ""]
for addr in addrs:
    da, dc = pts_a.get(addr), pts_c.get(addr)
    if da is None:
        only_c += 1
        kind = "only-C"
    elif dc is None:
        only_a += 1
        kind = "only-A"
    elif da == dc:
        same += 1
        kind = "same"
    else:
        diff += 1
        kind = "diff"
    rows.append(f"{kind}\t{addr}\t{da or ''}\t{dc or ''}")

lines.append(f"A({sheet_a})={len(pts_a)}  C(PF-260)={len(pts_c)}  same={same} diff={diff} onlyA={only_a} onlyC={only_c}")
lines.append("")
lines.extend(rows)
(out / "pf-io-compare.txt").write_text("\n".join(lines), encoding="utf-8")
print("A", len(pts_a), "C", len(pts_c), "same", same, "diff", diff, "onlyA", only_a, "onlyC", only_c)
