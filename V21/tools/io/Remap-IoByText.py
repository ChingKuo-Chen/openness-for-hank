# -*- coding: utf-8 -*-
"""Match old Hardware tags to C-table points by description, then move names."""
from __future__ import annotations

import csv
import re
import xml.etree.ElementTree as ET
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
LEGACY = Path(r"C:\Users\Hank\Documents\Codex\2026-08-14\new-chat-3\outputs\TiaOpennessCheck")
ROOT = LEGACY  # Practice/HmiExport 仍在舊包；下一波再遷
OUT = ROOT / "Practice" / "io-merge"
IMPORT = OUT / "import"
HW_ROOT = ROOT / "HmiExport" / "Templates" / "26037_Cusor"
C_DIR = Path(r"C:\Users\Hank\Desktop\25017")
C_XLSX = next(p for p in C_DIR.glob("*.xlsx") if "63B" in p.name and p.name.endswith("新.xlsx"))
B_XLSX = Path(r"C:\Users\Hank\Desktop\24137ST_拉梅莎_630mm-37B 框絞機生產線.xlsx")
A_XLSX = Path(r"C:\Users\Hank\Desktop\22021ST_大亞_630mm-61B框絞機生產線.xlsx")

ADDR_CELL = re.compile(r"^(?:%?\s*)(I|Q|IW|QW|ID|QD|AIW|AQW)(\d+(?:\.\d+)?)$", re.I)
# Keep OP1/OP2 so the two panels do not collapse into one match.
STRIP = re.compile(
    r"LINE\s*|全線\s*|PBL|PB\s*|JOYSTICK|\(NO\)|\(NC\)|\(RLY\)|\(MC\)|"
    r"備用|Spare|預留|预留|\s+|　|[()（）,_、\-–]",
    re.I,
)
OP_RE = re.compile(r"OP\s*([123])", re.I)
DIR_PAIRS = (
    ("正", "逆"),
    ("升", "降"),
    ("頂", "退"),
    ("啟動", "停止"),
    ("a側", "b側"),
)

# plc, TIA export folder, C sheet, old Excel (B=24137 / A=22021), old sheet
JOBS = [
    ("25017_Main_PLC", "PLC_24137_Main_PLC", "LINE PLC", "B", "LINE PLC"),
    ("25017_TF_PLC", "PLC_24137_TF_PLC", "TF-260 PLC", "B", "TF PLC"),
    ("25017_6B_Inside_PLC", "PLC_24137_6B_Inside_PLC", "Rotor inside PLC", "B", "Rotor inside PLC"),
    ("25017_12B_Inside_PLC", "PLC_24137_12B_Inside_PLC", "Rotor inside PLC", "B", "Rotor inside PLC"),
    ("25017_18B_Inside_PLC", "PLC_24137_18B_Inside_PLC", "Rotor inside PLC", "B", "Rotor inside PLC"),
    ("25017_PF_PLC", "PLC_25017_PF_PLC", "PF-260 PLC", "A", "給線PLC"),
]


def cell_str(v) -> str:
    if v is None or hasattr(v, "isoformat"):
        return ""
    return str(v).replace("\n", " ").strip()


def norm_addr(raw: str):
    m = ADDR_CELL.match(raw.replace(" ", ""))
    if not m:
        return None
    kind, num = m.group(1).upper(), m.group(2)
    if kind == "AIW":
        kind = "IW"
    if kind == "AQW":
        kind = "QW"
    return f"%{kind}{num}"


def looks_desc(s: str) -> bool:
    if not s or len(s) < 2:
        return False
    if ADDR_CELL.match(s.replace(" ", "")):
        return False
    return s.lower() not in {"hub", "x1", "p1", "p2", "no.", "no"}


def parse_sheet(ws):
    by = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        for ci, v in enumerate(row):
            addr = norm_addr(cell_str(v))
            if not addr:
                continue
            desc = ""
            if ci + 1 < len(row) and looks_desc(cell_str(row[ci + 1])):
                desc = cell_str(row[ci + 1])
            prev = by.get(addr)
            if prev is None or len(desc) > len(prev):
                by[addr] = desc
    return by


def parse_hw(path: Path):
    tags = []
    tree = ET.parse(path)
    for tag in tree.getroot().iter():
        if not tag.tag.endswith("PlcTag"):
            continue
        attrs, comment = {}, ""
        for child in list(tag):
            local = child.tag.split("}")[-1]
            if local == "AttributeList":
                for a in list(child):
                    attrs[a.tag.split("}")[-1]] = (a.text or "").strip()
            elif local == "ObjectList":
                for text_el in child.iter():
                    if text_el.tag.split("}")[-1] == "Text" and (text_el.text or "").strip():
                        comment = text_el.text.strip()
                        break
        if attrs.get("LogicalAddress") and attrs.get("Name"):
            tags.append({
                "name": attrs["Name"],
                "address": attrs["LogicalAddress"],
                "type": attrs.get("DataTypeName", "Bool"),
                "comment": comment,
            })
    return tags, tree


def fold(s: str) -> str:
    return STRIP.sub("", s or "").casefold()


def is_spare(desc: str) -> bool:
    if not (desc or "").strip():
        return True
    return bool(re.search(r"spare|備用|预留|預留", desc, re.I))


def io_kind(addr: str) -> str:
    body = addr[1:]
    if body.startswith("IW"):
        return "IW"
    if body.startswith("QW"):
        return "QW"
    if body.startswith("ID"):
        return "ID"
    if body.startswith("QD"):
        return "QD"
    if body.startswith("I"):
        return "I"
    if body.startswith("Q"):
        return "Q"
    return "?"


def op_num(s: str):
    m = OP_RE.search(s or "")
    return m.group(1) if m else None


def dir_conflict(a: str, b: str) -> bool:
    fa, fb = fold(a), fold(b)
    for left, right in DIR_PAIRS:
        if (left in fa and right in fb and left not in fb) or (
            right in fa and left in fb and right not in fb
        ):
            return True
    return False


NAME_ALIASES = (
    (r"(?i)\bcap\b", "引取"),
    (r"(?i)18b", "20b"),
    (r"(?i)ups\s*bit\s*0", "ups狀態1"),
    (r"(?i)ups\s*bit\s*1", "ups狀態2"),
    (r"(?i)ups\s*bit\s*2", "ups狀態3"),
    (r"(?i)main\s*power", "停電檢出"),
)


def name_as_text(name: str) -> str:
    s = re.sub(r"^[IQ]_[SF]?_", "", name or "")
    s = s.replace("_", " ")
    for pat, repl in NAME_ALIASES:
        s = re.sub(pat, repl, s)
    return s


def ratio(a: str, b: str) -> float:
    fa, fb = fold(a), fold(b)
    if not fa and not fb:
        return 1.0
    if not fa or not fb:
        return 0.0
    return SequenceMatcher(None, fa, fb).ratio()


def tag_texts(t: dict) -> list[str]:
    return [
        t.get("comment") or "",
        name_as_text(t.get("name") or ""),
        t.get("old_excel") or "",
    ]


def tag_ratio(t: dict, desc: str) -> float:
    return max(ratio(s, desc) for s in tag_texts(t))


def hex_id(n: int) -> str:
    return format(n, "X")


def make_tag(tid: int, name: str, addr: str, dtype: str, comment: str) -> ET.Element:
    tag = ET.Element("SW.Tags.PlcTag", ID=hex_id(tid), CompositionName="Tags")
    al = ET.SubElement(tag, "AttributeList")
    ET.SubElement(al, "DataTypeName").text = dtype
    ET.SubElement(al, "ExternalAccessible").text = "true"
    ET.SubElement(al, "ExternalVisible").text = "true"
    ET.SubElement(al, "ExternalWritable").text = "true"
    ET.SubElement(al, "LogicalAddress").text = addr
    ET.SubElement(al, "Name").text = name
    ol = ET.SubElement(tag, "ObjectList")
    mt = ET.SubElement(ol, "MultilingualText", ID=hex_id(tid + 1), CompositionName="Comment")
    items = ET.SubElement(mt, "ObjectList")
    for i, cult in enumerate(("en-US", "zh-CN", "zh-TW")):
        item = ET.SubElement(items, "MultilingualTextItem", ID=hex_id(tid + 2 + i), CompositionName="Items")
        ial = ET.SubElement(item, "AttributeList")
        ET.SubElement(ial, "Culture").text = cult
        ET.SubElement(ial, "Text").text = comment or ""
    return tag


def write_xml(src: Path, tags: list[dict], dest: Path):
    tree = ET.parse(src)
    root = tree.getroot()
    table = next(el for el in root.iter() if el.tag.endswith("PlcTagTable"))
    obj = next(ch for ch in list(table) if ch.tag.endswith("ObjectList"))
    for ch in list(obj):
        obj.remove(ch)
    tid = 1
    for t in tags:
        obj.append(make_tag(tid, t["name"], t["address"], t["type"], t["comment"]))
        tid += 6
    dest.parent.mkdir(parents=True, exist_ok=True)
    tree.write(dest, encoding="utf-8", xml_declaration=True)


def remap_plc(plc: str, export_folder: str, excel: dict, old_excel: dict) -> dict:
    src = HW_ROOT / export_folder / "TagTables" / "Hardware.xml"
    old_tags, _ = parse_hw(src)
    for t in old_tags:
        t["old_excel"] = old_excel.get(t["address"], "")

    usable_old = [t for t in old_tags if t["name"] and not is_spare(t["name"])]
    usable_c = {a: d for a, d in excel.items() if not is_spare(d)}

    candidates = []
    for t in usable_old:
        for addr, desc in usable_c.items():
            if io_kind(t["address"]) != io_kind(addr):
                continue
            if any(dir_conflict(s, desc) for s in tag_texts(t) if s):
                continue
            old_op, new_op = op_num(t["comment"] or t["old_excel"]), op_num(desc)
            if old_op and new_op and old_op != new_op:
                continue
            r = tag_ratio(t, desc)
            if r < 0.55:
                continue
            bonus = 0.05 if old_op and new_op and old_op == new_op else 0.0
            candidates.append((r + bonus, r, t, addr, desc))

    candidates.sort(key=lambda x: (-x[0], x[2]["name"], x[3]))
    used_old = set()
    used_addr = set()
    moved, same_addr, fuzzy = [], [], []

    for score, r, t, addr, desc in candidates:
        if t["name"] in used_old or addr in used_addr:
            continue
        used_old.add(t["name"])
        used_addr.add(addr)
        item = {
            "name": t["name"],
            "type": t["type"],
            "old_address": t["address"],
            "address": addr,
            "old_comment": t["comment"] or t.get("old_excel", ""),
            "comment": desc,
            "ratio": round(r, 2),
            "kind": "moved" if t["address"] != addr else "same-addr",
        }
        if r >= 0.80:
            (moved if t["address"] != addr else same_addr).append(item)
        else:
            item["kind"] = "fuzzy"
            fuzzy.append(item)

    unmatched_c = [
        {"address": a, "comment": d}
        for a, d in excel.items()
        if a not in used_addr and not is_spare(d)
    ]
    unmatched_spare = [
        {"address": a, "comment": d}
        for a, d in excel.items()
        if a not in used_addr and is_spare(d)
    ]

    leftover = []
    taken = {i["address"] for i in moved + same_addr + fuzzy}
    for t in old_tags:
        if t["name"] in used_old:
            continue
        item = {
            "name": t["name"],
            "type": t["type"],
            "old_address": t["address"],
            "address": t["address"],
            "old_comment": t["comment"] or t.get("old_excel", ""),
            "comment": t["comment"] or t.get("old_excel", ""),
            "kind": "old-unmatched",
        }
        if t["address"] in taken:
            item["kind"] = "displaced"
        leftover.append(item)

    merged = []
    for group in (moved, same_addr, fuzzy):
        merged.extend(group)
    for t in leftover:
        if t["kind"] != "displaced":
            merged.append(t)

    dest = IMPORT / f"PLC_{plc}" / "TagTables" / "Hardware.xml"
    write_xml(src, merged, dest)
    return {
        "plc": plc,
        "moved": moved,
        "same_addr": same_addr,
        "fuzzy": fuzzy,
        "unmatched_c": unmatched_c,
        "unmatched_spare": unmatched_spare,
        "leftover": leftover,
        "merged": merged,
    }


def load_book(path: Path) -> dict[str, dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {sn: parse_sheet(wb[sn]) for sn in wb.sheetnames}
    wb.close()
    return sheets


def main():
    sheets = load_book(C_XLSX)
    old_books = {"B": load_book(B_XLSX), "A": load_book(A_XLSX)}

    results = []
    for plc, folder, sheet, old_kind, old_sheet in JOBS:
        excel = sheets.get(sheet, {})
        old_excel = old_books[old_kind].get(old_sheet, {})
        results.append(remap_plc(plc, folder, excel, old_excel))

    lines = [
        "# IO 依文字重對（舊 Tag 名 → C 表新位址）",
        "",
        "對到的：舊名字跟著功能搬到新位址。對不到的 C 點可能是新加的，先列出來不自動建。",
        "",
    ]
    for r in results:
        lines.append(f"## {r['plc']}")
        lines.append(
            f"- 搬走 {len(r['moved'])}、同位址仍對上 {len(r['same_addr'])}、模糊 {len(r['fuzzy'])}"
        )
        lines.append(
            f"- C 對不到（可能新加）{len(r['unmatched_c'])}、C 備用 {len(r['unmatched_spare'])}"
        )
        lines.append(
            f"- 舊點對不到 {sum(1 for x in r['leftover'] if x['kind']=='old-unmatched')}、"
            f"舊點被擠走 {sum(1 for x in r['leftover'] if x['kind']=='displaced')}"
        )
        if r["moved"]:
            lines.append("- **搬走（舊名到新址）：**")
            for x in r["moved"]:
                lines.append(
                    f"  - `{x['name']}` {x['old_address']} → {x['address']}  "
                    f"{x['old_comment'][:28]} ⇒ {x['comment'][:28]}  ({x['ratio']:.0%})"
                )
        if r["fuzzy"]:
            lines.append("- **模糊（55–80%，已搬或留）：**")
            for x in r["fuzzy"]:
                arrow = f"{x['old_address']} → {x['address']}" if x["old_address"] != x["address"] else x["address"]
                lines.append(
                    f"  - `{x['name']}` {arrow}  {x['old_comment'][:28]} ⇒ {x['comment'][:28]}  ({x['ratio']:.0%})"
                )
        if r["unmatched_c"]:
            lines.append("- **C 對不到，可能是新 IO（先不建，要討論）：**")
            for x in r["unmatched_c"]:
                lines.append(f"  - {x['address']}  {x['comment']}")
        displaced = [x for x in r["leftover"] if x["kind"] == "displaced"]
        if displaced:
            lines.append("- **舊名字程式可能還在用，但新表沒對到且舊位址已被佔：**")
            for x in displaced:
                lines.append(f"  - `{x['name']}` 原 {x['old_address']}  {x['old_comment'][:40]}")
        lines.append("")

    report = OUT / "remap-listing.txt"
    report.write_text("\n".join(lines), encoding="utf-8")

    with (OUT / "remap.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(
            f, fieldnames=["plc", "kind", "name", "old_address", "address", "old_comment", "comment", "ratio"]
        )
        w.writeheader()
        for r in results:
            for x in r["moved"] + r["same_addr"] + r["fuzzy"]:
                w.writerow({"plc": r["plc"], **{k: x.get(k, "") for k in w.fieldnames if k != "plc"}})
            for x in r["unmatched_c"]:
                w.writerow({
                    "plc": r["plc"], "kind": "new-c", "name": "", "old_address": "",
                    "address": x["address"], "old_comment": "", "comment": x["comment"], "ratio": "",
                })
            for x in r["leftover"]:
                w.writerow({"plc": r["plc"], **{k: x.get(k, "") for k in w.fieldnames if k != "plc"}})

    print("C", C_XLSX)
    print("report", report)
    for r in results:
        print(
            r["plc"],
            "moved", len(r["moved"]),
            "same", len(r["same_addr"]),
            "fuzzy", len(r["fuzzy"]),
            "newC", len(r["unmatched_c"]),
            "displaced", sum(1 for x in r["leftover"] if x["kind"] == "displaced"),
        )


if __name__ == "__main__":
    main()
